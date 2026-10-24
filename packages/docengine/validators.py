"""Post-render validation: re-open every artifact and check it before anything is delivered.

Returns a list of human-readable problems (empty list == valid). The thinker retries/escalates
on problems instead of sending a broken file.
"""
from __future__ import annotations

import csv
import io
import json
import re
import unicodedata
import zipfile
from pathlib import Path
from typing import Iterable, List, Optional

from .model import Document, Table
from .renderers import RenderedFile

FONT_DIRS = ("/usr/share/fonts", "/usr/local/share/fonts", str(Path.home() / ".fonts"),
             "/Library/Fonts", "/System/Library/Fonts")
_IGNORABLE = {"Cc", "Cf", "Zs", "Zl", "Zp"}  # controls, ZWJ/ZWNJ (format), spaces


def find_font_file(family: str) -> Optional[Path]:
    """'Noto Sans Devanagari' -> .../NotoSansDevanagari-Regular.ttf (first match), else None."""
    stem = family.replace(" ", "")
    for d in FONT_DIRS:
        root = Path(d)
        if not root.exists():
            continue
        for pattern in (f"{stem}-Regular.*", f"{stem}.*", f"{stem}*Regular*.*"):
            hits = sorted(root.rglob(pattern))
            if hits:
                return hits[0]
    return None


def missing_glyphs(text: str, font_path: Path) -> List[str]:
    """Characters in `text` that the font's cmap cannot render (would show as tofu)."""
    from fontTools.ttLib import TTCollection, TTFont

    fonts = TTCollection(str(font_path)).fonts if str(font_path).endswith(".ttc") else [TTFont(str(font_path))]
    cmap: set = set()
    for f in fonts:
        cmap |= set(f.getBestCmap() or {})
    return sorted({ch for ch in text if unicodedata.category(ch) not in _IGNORABLE and ord(ch) not in cmap})


def _non_latin(text: str) -> bool:
    return any(ord(ch) >= 0x0590 for ch in text)  # Arabic/Urdu (U+0600) and Indic (U+0900+)


def _ooxml_order_problems(data: bytes) -> List[str]:
    """Word refuses files whose rPr/pPr children are out of schema order; LibreOffice silently tolerates them."""
    import zipfile

    from docx.oxml.ns import qn
    from lxml import etree

    from .renderers.docx_r import PPR_SEQ, RPR_SEQ

    problems: List[str] = []
    seqs = {qn("w:rPr"): RPR_SEQ, qn("w:pPr"): PPR_SEQ}
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        for part in ("word/document.xml", "word/styles.xml"):
            root = etree.fromstring(z.read(part))
            for parent_tag, seq in seqs.items():
                pos = {qn(t): i for i, t in enumerate(seq)}
                for el in root.iter(parent_tag):
                    idx = [pos[c.tag] for c in el if c.tag in pos]
                    if idx != sorted(idx):
                        problems.append(f"docx: out-of-order children in {etree.QName(parent_tag).localname} of {part}")
                        break
    return problems


def _docx(data: bytes, doc: Document) -> List[str]:
    from docx import Document as Docx
    from docx.oxml.ns import qn

    problems: List[str] = []
    d = Docx(io.BytesIO(data))
    text = "\n".join(p.text for p in d.paragraphs)
    if doc.title not in text:
        problems.append("docx: title missing")
    for p in d.paragraphs:
        for r in p.runs:
            if not _non_latin(r.text):
                continue
            rpr = r._element.rPr
            fonts = rpr.find(qn("w:rFonts")) if rpr is not None else None
            if fonts is None or not fonts.get(qn("w:cs")):
                problems.append(f"docx: complex-script font not set on run {r.text[:20]!r}")
            if r.bold and rpr.find(qn("w:bCs")) is None:
                problems.append(f"docx: bold not applied to complex script in {r.text[:20]!r}")
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        for name in z.namelist():
            if re.fullmatch(r"word/charts/chart\d+\.xml", name) and b"externalData" not in z.read(name):
                problems.append(f"docx: {name} has no workbook (Apple previews hang)")
    problems.extend(_ooxml_order_problems(data))
    return problems


def _xlsx(data: bytes, doc: Document) -> List[str]:
    from openpyxl import load_workbook

    problems: List[str] = []
    wb = load_workbook(io.BytesIO(data))  # formulas preserved (not data_only)
    tables = [b for b in doc.blocks if isinstance(b, Table)]
    if len(wb.sheetnames) < len(tables):
        problems.append(f"xlsx: expected ≥{len(tables)} sheets, got {len(wb.sheetnames)}")
    for ws, t in zip(wb.worksheets, tables):
        header = [c.value for c in ws[1]][: len(t.columns)]
        if header != [c.name for c in t.columns]:
            problems.append(f"xlsx: header mismatch in sheet {ws.title!r}")
        if t.total_row and any(c.type in ("number", "currency") for c in t.columns) and t.rows:
            total_row = ws[2 + len(t.rows)]
            if not any(isinstance(c.value, str) and c.value.startswith("=SUM(") for c in total_row):
                problems.append(f"xlsx: total row has no SUM formula in sheet {ws.title!r}")
    return problems


def _pptx(data: bytes, doc: Document) -> List[str]:
    from pptx import Presentation
    from pptx.oxml.ns import qn

    prs = Presentation(io.BytesIO(data))
    if len(prs.slides) == 0:
        return ["pptx: no slides"]
    problems = []
    first = prs.slides[0].shapes.title
    if first is None or first.text != doc.title:
        problems.append("pptx: title slide missing title")
    has_notes = any(r.reltype.endswith("/notesMaster") for r in prs.part.rels.values())
    if has_notes and prs.part._element.find(qn("p:notesMasterIdLst")) is None:
        problems.append("pptx: notes master not listed (Apple previews hang)")
    return problems


def _font_descriptor(font):
    desc = font.get("/FontDescriptor")
    if desc is None and font.get("/Subtype") == "/Type0":  # CID fonts keep the descriptor on the descendant
        descendants = font.get("/DescendantFonts")
        if descendants:
            desc = descendants[0].get_object().get("/FontDescriptor")
    return desc.get_object() if desc is not None else None


def _iter_fonts(resources, seen):
    """Yield font dicts from a resources dict, recursing into form XObjects (cycle-safe)."""
    if resources is None:
        return
    resources = resources.get_object()
    for ref in (resources.get("/Font") or {}).values():
        yield ref.get_object()
    for ref in (resources.get("/XObject") or {}).values():
        xobj = ref.get_object()
        key = id(xobj)
        if xobj.get("/Subtype") == "/Form" and key not in seen:
            seen.add(key)
            yield from _iter_fonts(xobj.get("/Resources"), seen)


def _pdf(data: bytes, doc: Document) -> List[str]:
    """Structural check (fonts may sit in compressed object streams, so never grep the raw bytes)."""
    if not data.startswith(b"%PDF"):
        return ["pdf: not a PDF"]
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if len(reader.pages) == 0:
        return ["pdf: no pages"]
    problems = set()
    for page in reader.pages:
        for font in _iter_fonts(page.get("/Resources"), set()):
            if font.get("/Subtype") == "/Type3":
                continue  # glyphs are drawn inline; nothing to embed
            desc = _font_descriptor(font)
            if desc is None or not any(k in desc for k in ("/FontFile", "/FontFile2", "/FontFile3")):
                problems.add(f"pdf: font {font.get('/BaseFont')} not embedded")
    return sorted(problems)


def _textual(data: bytes, doc: Document, fmt: str) -> List[str]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return [f"{fmt}: not valid UTF-8"]
    if fmt == "json":
        try:
            Document.model_validate(json.loads(text))
        except Exception as e:
            return [f"json: does not round-trip: {e}"]
        return []
    if fmt == "csv":
        rows = list(csv.reader(io.StringIO(text)))
        return [] if rows else ["csv: empty"]
    title = doc.title if fmt != "html" else doc.title.replace("&", "&amp;")
    return [] if title in text else [f"{fmt}: title missing"]


def validate(doc: Document, rendered: RenderedFile) -> List[str]:
    fmt = rendered.fmt
    if not rendered.data:
        return [f"{fmt}: empty file"]
    checks = {"docx": _docx, "xlsx": _xlsx, "pptx": _pptx, "pdf": _pdf}
    if fmt in checks:
        return checks[fmt](rendered.data, doc)
    return _textual(rendered.data, doc, fmt)


def glyph_report(doc: Document, families: Iterable[str]) -> dict:
    """For each available font family, which characters of the document it cannot render."""
    text = doc.all_text()
    report = {}
    for fam in families:
        path = find_font_file(fam)
        report[fam] = None if path is None else missing_glyphs(text, path)
    return report
