"""XLSX renderer (openpyxl): one sheet per table (real SUM formulas, Indian number formats), a sheet per chart (its
numbers and an Excel chart drawn from them), a Summary that links each table's totals, prose -> Notes. The table
sheets come first, in order (validators.py checks them by position)."""
from __future__ import annotations

import io
import re
from datetime import datetime
from typing import List, Optional

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter, quote_sheetname

from ..charts import PALETTES, as_table
from ..i18n import XLSX_FORMATS, meta_label, parse_number, whole_rupees
from ..model import Chart, Document, Heading, KeyValue, ListBlock, Paragraph, Signature, Slide, Table

HEADER_FILL = PatternFill("solid", fgColor="E7E5E4")
_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
_TABLE_NO = re.compile(r"^\s*(?:table|तालिका|सारणी)\s*[\d.]+\s*[:.\-–]?\s*", re.I)  # "Table 2.1: "


def _sheet_name(raw: str, used: set) -> str:
    """A tab's name (31 characters at most): "Table 2.1: Three-year financial projections" -> "Three-year financial
    projections" is cut at a word, not mid-word."""
    text = _TABLE_NO.sub("", raw).strip() or raw
    text = " ".join(_BAD_SHEET_CHARS.sub(" ", text).split())
    if len(text) > 31:
        cut = text[:31]
        text = (cut[:cut.rfind(" ")] if " " in cut[10:] else cut).rstrip(" ,;-–(")
    base = text or "Sheet"
    name, n = base, 2
    while name.lower() in used:
        suffix = f" ({n})"
        name = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(name.lower())
    return name


def _to_date(value) -> Optional[datetime]:
    if not isinstance(value, str):
        return None
    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%d %B %Y", "%d %b %Y"):
        try:
            return datetime.strptime(value.strip(), fmt)
        except ValueError:
            continue
    return None


def _fmt(t: Table, ci: int, whole: List[bool]) -> Optional[str]:
    """Column ci's number format: whole rupees where every amount is whole (i18n.whole_rupees)."""
    return XLSX_FORMATS.get("currency_whole" if whole[ci] else t.columns[ci].type)


def _write_table(ws, t: Table, rtl: bool, total_word: str = "Total", doc_type: str = "report") -> Optional[int]:
    """Returns the row of the total (SUM) row, if there is one."""
    whole = whole_rupees(t, doc_type)
    ws.append([c.name for c in t.columns])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in t.rows:
        out: List[object] = []
        for ci, val in enumerate(row):
            ctype = t.columns[ci].type
            if ctype in ("number", "currency", "percent"):
                num = parse_number(val)
                out.append(num if num is not None else val)
            elif ctype == "date":
                out.append(_to_date(val) or val)
            else:
                out.append(val)
        ws.append(out)
    last_data_row = 1 + len(t.rows)
    for ci, _col in enumerate(t.columns, start=1):
        fmt = _fmt(t, ci - 1, whole)
        if fmt:
            for r in range(2, last_data_row + 1):
                ws.cell(row=r, column=ci).number_format = fmt
    numeric = [c.type in ("number", "currency") for c in t.columns]
    total_row = None
    if t.total_row and any(numeric) and t.rows:
        total_row = last_data_row + 1
        first_text = next((i for i, n in enumerate(numeric) if not n), None)
        if first_text is not None:
            ws.cell(row=total_row, column=first_text + 1, value=total_word).font = Font(bold=True)
        for ci, is_num in enumerate(numeric, start=1):
            if is_num:
                letter = get_column_letter(ci)
                c = ws.cell(row=total_row, column=ci, value=f"=SUM({letter}2:{letter}{last_data_row})")
                c.font = Font(bold=True)
                c.number_format = _fmt(t, ci - 1, whole)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(t.columns))}{last_data_row}"
    for ci, col in enumerate(t.columns, start=1):
        width = max([len(col.name)] + [len(str(r[ci - 1])) for r in t.rows if r[ci - 1] is not None] + [8])
        ws.column_dimensions[get_column_letter(ci)].width = min(width + 4, 50)
    ws.sheet_view.rightToLeft = rtl
    return total_row


def _write_chart(ws, c: Chart, theme: str, rtl: bool) -> None:
    """The chart's numbers as a table, and an Excel chart drawn from those cells (edit a number, the chart follows)."""
    _write_table(ws, as_table(c), rtl)
    n_rows, n_cols = len(c.categories), len(c.series)
    cats = Reference(ws, min_col=1, min_row=2, max_row=1 + n_rows)
    data = Reference(ws, min_col=2, max_col=1 + n_cols, min_row=1, max_row=1 + n_rows)
    chart = {"pie": PieChart, "line": LineChart}.get(c.kind, BarChart)()
    if c.kind in ("bar", "column"):
        chart.type = "bar" if c.kind == "bar" else "col"
        chart.grouping = "clustered"
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.title = c.title
    chart.height, chart.width = 9, 16
    palette = PALETTES.get(theme, PALETTES["boldoc"])
    if c.kind != "pie":
        for i, s in enumerate(chart.series):
            color = palette[i % len(palette)]
            if c.kind == "line":
                s.graphicalProperties.line.solidFill = color
                s.graphicalProperties.line.width = 28000
                s.smooth = False
            else:
                s.graphicalProperties.solidFill = color
                s.graphicalProperties.line.solidFill = color
    ws.add_chart(chart, f"{get_column_letter(n_cols + 3)}2")


def _write_summary(ws, rows: List[tuple]) -> None:
    """Each table sheet's totals, as formulas pointing at its Total row (change a sheet, the summary follows)."""
    ws.append(["Sheet", "What", "Total"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for sheet, column, ref, fmt in rows:
        ws.append([sheet, column, f"={ref}"])
        ws.cell(row=ws.max_row, column=3).number_format = fmt
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width, ws.column_dimensions["C"].width = 28, 28, 20
    ws.freeze_panes = "A2"


def _prose_lines(doc: Document) -> List[tuple]:
    lines: List[tuple] = []
    for k, v in doc.meta.items():
        lines.append((f"{k.capitalize()}: {v}", False))
    for b in doc.blocks:
        if isinstance(b, Heading):
            lines.append((b.text, True))
        elif isinstance(b, Paragraph):
            lines.append((b.text, False))
        elif isinstance(b, ListBlock):
            lines.extend(((f"{i + 1}. " if b.ordered else "• ") + it, False) for i, it in enumerate(b.items))
        elif isinstance(b, KeyValue):
            lines.extend((f"{p[0]}: {p[1]}", False) for p in b.pairs if len(p) >= 2)
        elif isinstance(b, Signature):
            lines.extend([(b.closing, False), (b.name, True)] + ([(b.designation, False)] if b.designation else []))
        elif isinstance(b, Slide):
            lines.append((b.title, True))
            lines.extend((f"• {x}", False) for x in b.bullets)
    return lines


def render(doc: Document) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    used: set = set()
    tables = [b for b in doc.blocks if isinstance(b, Table)]
    summary: List[tuple] = []
    for i, t in enumerate(tables, start=1):
        ws = wb.create_sheet(_sheet_name(t.caption or f"Table {i}", used))
        total_row = _write_table(ws, t, doc.rtl, meta_label("total", doc.lang), doc.doc_type)
        if total_row:
            whole = whole_rupees(t, doc.doc_type)
            for ci, col in enumerate(t.columns, start=1):
                if col.type in ("number", "currency"):
                    ref = f"{quote_sheetname(ws.title)}!{get_column_letter(ci)}{total_row}"
                    summary.append((ws.title, col.name, ref, _fmt(t, ci - 1, whole)))
    for i, c in enumerate([b for b in doc.blocks if isinstance(b, Chart)], start=1):
        _write_chart(wb.create_sheet(_sheet_name(c.title or f"Chart {i}", used)), c, doc.style.theme, doc.rtl)
    if len(tables) >= 2 and summary:
        _write_summary(wb.create_sheet(_sheet_name("Summary", used)), summary)
    prose = _prose_lines(doc)
    if prose or not tables:
        ws = wb.create_sheet(_sheet_name("Notes" if tables else doc.title or "Document", used))
        ws.append([doc.title])
        ws["A1"].font = Font(bold=True, size=14)
        for text, bold in prose:
            ws.append([text])
            ws.cell(row=ws.max_row, column=1).font = Font(bold=bold)
        ws.column_dimensions["A"].width = 100
        for row in ws.iter_rows():
            row[0].alignment = Alignment(wrap_text=True, vertical="top")
        ws.sheet_view.rightToLeft = doc.rtl
    wb.properties.title = doc.title
    wb.properties.creator = "Sarvam Doc Agent"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
