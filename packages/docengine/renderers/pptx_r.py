"""PPTX renderer (python-pptx): 16:9 decks in a theme (style.theme: boldoc, plain or dark).

Explicit Slide blocks keep their layout (bullets, title, section, two_column, quote, closing); a table or chart block
gets a slide of its own, never dropped; a heading in a deck is a section divider. Without Slide blocks a deck is
derived: a title slide, a slide per heading section (paragraphs and list items become bullets), a slide per table and
chart. Long bullet lists and tables continue on the next slide. Every slide keeps a real title (outline view, screen
readers) on the title-only layout; the rest is drawn in the theme. Complex scripts need the a:cs typeface on every
run; RTL needs a:pPr/@rtl.
"""
from __future__ import annotations

import io
from typing import List, Optional

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from ..charts import ooxml_data, ooxml_type, style_native
from ..deck import SLIDE_BULLETS, SLIDE_ROWS, subtitle, theme
from ..i18n import fonts_for, format_cell, meta_label, parse_number, whole_rupees
from ..model import Chart, Document, Heading, KeyValue, ListBlock, Paragraph, Signature, Slide, Table

W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.7)


def _rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_)


class _Deck:
    def __init__(self, doc: Document):
        self.doc = doc
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self.t = theme(doc)
        f = fonts_for(doc.script)
        self.latin, self.cs, self.rtl = fonts_for("Latn")["office"], f["office"], doc.rtl
        self.n = 0

    # ------------------------------------------------------------------ text
    def fonts(self, text_frame, size: Optional[int] = None, color: Optional[str] = None, bold: Optional[bool] = None):
        for p in text_frame.paragraphs:
            if self.rtl:
                p._p.get_or_add_pPr().set("rtl", "1")
                p.alignment = PP_ALIGN.RIGHT
            for r in p.runs:
                r.font.name = self.latin  # creates a:latin
                if size:
                    r.font.size = Pt(size)
                if color:
                    r.font.color.rgb = _rgb(color)
                if bold is not None:
                    r.font.bold = bold
                # schema order inside a:rPr is latin, ea, cs: put cs right after latin (or ea, if present)
                rpr = r._r.get_or_add_rPr()
                cs_el = rpr.find(qn("a:cs"))
                if cs_el is None:
                    cs_el = rpr.makeelement(qn("a:cs"), {})
                    anchor = rpr.find(qn("a:ea"))
                    (anchor if anchor is not None else rpr.find(qn("a:latin"))).addnext(cs_el)
                cs_el.set("typeface", self.cs)

    def textbox(self, slide, x, y, w, h, lines: List[str], size: int, color: str, bold: bool = False,
                align=None, bullets: bool = False, anchor=MSO_ANCHOR.TOP):
        tb = slide.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = 0  # text lines up with the title, the accent bar and tables
        if bullets:
            self.bullet_style(tf, size)
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.space_after = Pt(size * 0.6 if bullets else 4)
            if align is not None:
                p.alignment = align
        self.fonts(tf, size, color, bold)
        return tb

    def bullet_style(self, tf, size: int) -> None:
        """Real bullets with a hanging indent, so a wrapped line starts under the text, not under the bullet. Set on
        the text box's list style: Apple's Office importer misplaces bullets indented paragraph by paragraph."""
        lst = tf._txBody.find(qn("a:lstStyle"))
        indent = int(Pt(size) * 1.1)
        lvl = lst.makeelement(qn("a:lvl1pPr"), {"marL": str(indent), "indent": str(-indent)})
        clr = lvl.makeelement(qn("a:buClr"), {})
        clr.append(clr.makeelement(qn("a:srgbClr"), {"val": self.t["accent"]}))
        lvl.append(clr)  # schema order: bullet colour, bullet font, bullet char
        lvl.append(lvl.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
        lvl.append(lvl.makeelement(qn("a:buChar"), {"char": "•"}))
        lst.append(lvl)

    @staticmethod
    def rule(cell, edge: str, color: str) -> None:
        """A 1.5 pt rule on one edge of a table cell (a:lnT or a:lnB), as PowerPoint writes it."""
        tcpr = cell._tc.get_or_add_tcPr()
        ln = tcpr.makeelement(qn(f"a:ln{edge}"), {"w": str(Pt(1.5)), "cap": "flat", "cmpd": "sng", "algn": "ctr"})
        fill = ln.makeelement(qn("a:solidFill"), {})
        fill.append(fill.makeelement(qn("a:srgbClr"), {"val": color}))
        ln.append(fill)
        ln.append(ln.makeelement(qn("a:prstDash"), {"val": "solid"}))
        # schema order in a:tcPr: lnL, lnR, lnT, lnB, ..., then the fill
        anchor = tcpr.find(qn("a:lnT")) if edge == "B" else None
        if anchor is not None:
            anchor.addnext(ln)
        else:
            tcpr.insert(0, ln)

    # ------------------------------------------------------------------ frames
    def slide(self, title: str, layout: int = 5, chrome: bool = True):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[layout])
        self.n += 1
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = _rgb(self.t["bg"])
        if s.shapes.title is not None:
            ttl = s.shapes.title
            ttl.left, ttl.top, ttl.width, ttl.height = MARGIN, Inches(0.45), W - 2 * MARGIN, Inches(1.0)
            ttl.text = title
            ttl.text_frame.word_wrap = True
            ttl.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            ttl.text_frame.margin_left = ttl.text_frame.margin_right = 0
            for p in ttl.text_frame.paragraphs:
                p.alignment = PP_ALIGN.RIGHT if self.rtl else PP_ALIGN.LEFT
            self.fonts(ttl.text_frame, 30 if len(title) < 50 else 24, self.t["title"], True)
        if chrome:
            bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, MARGIN, Inches(1.5), Inches(1.1), Inches(0.06))
            bar.fill.solid()
            bar.fill.fore_color.rgb = _rgb(self.t["accent"])
            bar.line.fill.background()
            self.textbox(s, MARGIN, H - Inches(0.55), W - 2 * MARGIN, Inches(0.35),
                         [f"{self.doc.title}   ·   {self.n}"], 10, self.t["muted"],
                         align=PP_ALIGN.LEFT if self.rtl else PP_ALIGN.RIGHT)
        return s

    def notes(self, s, text: Optional[str]) -> None:
        if text:
            s.notes_slide.notes_text_frame.text = text

    # ------------------------------------------------------------------ slides
    def title_slide(self, title: str, subtitle: str) -> None:
        s = self.prs.slides.add_slide(self.prs.slide_layouts[0])
        self.n += 1
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = _rgb(self.t["bg"])
        band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.35), H)
        band.fill.solid()
        band.fill.fore_color.rgb = _rgb(self.t["accent"])
        band.line.fill.background()
        ttl = s.shapes.title
        ttl.left, ttl.top, ttl.width, ttl.height = Inches(1.1), Inches(2.3), W - Inches(2.2), Inches(1.8)
        ttl.text = title
        ttl.text_frame.word_wrap = True
        ttl.text_frame.vertical_anchor = MSO_ANCHOR.BOTTOM
        for p in ttl.text_frame.paragraphs:
            p.alignment = PP_ALIGN.RIGHT if self.rtl else PP_ALIGN.LEFT
        self.fonts(ttl.text_frame, 44 if len(title) < 40 else 34, self.t["title"], True)
        sub = s.placeholders[1]
        if subtitle:
            sub.left, sub.top, sub.width, sub.height = Inches(1.1), Inches(4.3), W - Inches(2.2), Inches(1.0)
            sub.text = subtitle
            for p in sub.text_frame.paragraphs:
                p.alignment = PP_ALIGN.RIGHT if self.rtl else PP_ALIGN.LEFT
            self.fonts(sub.text_frame, 20, self.t["muted"])
        else:  # an empty placeholder would show 'Click to add subtitle' while editing
            sub._element.getparent().remove(sub._element)

    def bullets(self, title: str, items: List[str], notes: Optional[str] = None) -> None:
        items = [i for i in items if i and i.strip()] or [""]
        chunks = [items[i:i + SLIDE_BULLETS] for i in range(0, len(items), SLIDE_BULLETS)]
        for n, chunk in enumerate(chunks):
            s = self.slide(title if n == 0 else f"{title} (contd.)")
            longest = max(len(x) for x in chunk)
            size = 24 if len(chunk) <= 4 and longest < 90 else 20 if longest < 160 else 16
            self.textbox(s, MARGIN, Inches(1.85), W - 2 * MARGIN, H - Inches(2.6), chunk, size, self.t["text"],
                         bullets=True)
            self.notes(s, notes if n == 0 else None)

    def two_column(self, title: str, left: List[str], right: List[str], notes: Optional[str]) -> None:
        s = self.slide(title)
        col = (W - 2 * MARGIN - Inches(0.6)) / 2
        for i, items in enumerate((left, right)):
            self.textbox(s, MARGIN + i * (col + Inches(0.6)), Inches(1.85), col, H - Inches(2.6), items[:SLIDE_BULLETS],
                         20, self.t["text"], bullets=True)
        rule = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, MARGIN + col + Inches(0.29), Inches(1.95), Inches(0.02),
                                  H - Inches(2.9))
        rule.fill.solid()
        rule.fill.fore_color.rgb = _rgb(self.t["rule"])
        rule.line.fill.background()
        self.notes(s, notes)

    def quote(self, title: str, line: str, notes: Optional[str]) -> None:
        s = self.slide(title)
        self.textbox(s, MARGIN + Inches(0.6), Inches(2.2), W - 2 * MARGIN - Inches(1.2), Inches(3.2), [line],
                     48 if len(line) < 20 else 32 if len(line) < 90 else 24, self.t["accent"] if len(line) < 20
                     else self.t["text"], bold=len(line) < 20, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        self.notes(s, notes)

    def section(self, title: str, notes: Optional[str] = None) -> None:
        s = self.slide(title, chrome=False)
        panel = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(2.6), W, Inches(2.3))
        panel.fill.solid()
        panel.fill.fore_color.rgb = _rgb(self.t["panel"])
        panel.line.fill.background()
        s.shapes._spTree.remove(panel._element)
        s.shapes._spTree.insert(2, panel._element)  # behind the title
        ttl = s.shapes.title  # title and bar centred in the panel; a second title line grows upwards
        ttl.top, ttl.height = Inches(2.9), Inches(1.0)
        ttl.text_frame.vertical_anchor = MSO_ANCHOR.BOTTOM
        self.fonts(ttl.text_frame, 36, self.t["title"], True)
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, MARGIN, Inches(4.1), Inches(1.1), Inches(0.06))
        bar.fill.solid()
        bar.fill.fore_color.rgb = _rgb(self.t["accent"])
        bar.line.fill.background()
        self.notes(s, notes)

    def closing(self, title: str, lines: List[str], notes: Optional[str]) -> None:
        s = self.slide(title, chrome=False)
        ttl = s.shapes.title
        ttl.top, ttl.height = Inches(2.4), Inches(1.3)
        for p in ttl.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
        self.fonts(ttl.text_frame, 40, self.t["title"], True)
        if lines:
            self.textbox(s, MARGIN, Inches(3.9), W - 2 * MARGIN, Inches(2.0), lines[:4], 20, self.t["muted"],
                         align=PP_ALIGN.CENTER)
        self.notes(s, notes)

    def table(self, t: Table, fallback_title: str) -> None:
        title = t.caption or fallback_title
        numeric = [c.type in ("number", "currency") for c in t.columns]
        whole = whole_rupees(t, self.doc.doc_type)
        add_total = t.total_row and any(numeric) and bool(t.rows)
        chunks = [t.rows[i:i + SLIDE_ROWS] for i in range(0, max(1, len(t.rows)), SLIDE_ROWS)] or [[]]
        for n, rows in enumerate(chunks):
            last = n == len(chunks) - 1
            total_here = add_total and last
            s = self.slide(title if n == 0 else f"{title} (contd.)")
            n_rows = 1 + len(rows) + int(total_here)
            row_h = Inches(0.42)
            shape = s.shapes.add_table(n_rows, len(t.columns), MARGIN, Inches(1.85), W - 2 * MARGIN,
                                       Emu(row_h * n_rows))
            tbl = shape.table
            size = 14 if len(t.columns) <= 4 else 12 if len(t.columns) <= 6 else 10
            for ci, col in enumerate(t.columns):
                cell = tbl.cell(0, ci)
                cell.text = col.name
                cell.fill.solid()
                cell.fill.fore_color.rgb = _rgb(self.t["accent"] if self.doc.style.theme != "plain" else "1F3A5F")
                self.fonts(cell.text_frame, size, "FFFFFF", True)
            for ri, row in enumerate(rows, start=1):
                for ci, val in enumerate(row):
                    cell = tbl.cell(ri, ci)
                    cell.text = format_cell(val, t.columns[ci].type, whole[ci])
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = _rgb(self.t["panel"] if ri % 2 == 0 else self.t["bg"])
                    if numeric[ci]:
                        cell.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
                    self.fonts(cell.text_frame, size, self.t["text"])
            if total_here:
                last_row = n_rows - 1
                first_text = next((i for i, num in enumerate(numeric) if not num), 0)
                for ci in range(len(t.columns)):
                    cell = tbl.cell(last_row, ci)
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = _rgb(self.t["panel"])
                    if ci == first_text:
                        cell.text = meta_label("total", self.doc.lang)
                    elif numeric[ci]:
                        cell.text = format_cell(sum(parse_number(r[ci]) or 0 for r in t.rows), t.columns[ci].type,
                                                whole[ci])
                        cell.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
                    self.fonts(cell.text_frame, size, self.t["text"], True)
                    self.rule(cell, "T", self.t["text"])  # the total reads as a total, not one more banded row
                    if last_row > 1:
                        self.rule(tbl.cell(last_row - 1, ci), "B", self.t["text"])

    def chart(self, c: Chart, fallback_title: str) -> None:
        s = self.slide(c.title or fallback_title)
        frame = s.shapes.add_chart(ooxml_type(c), MARGIN, Inches(1.8), W - 2 * MARGIN, H - Inches(2.55),
                                   ooxml_data(c))
        chart = frame.chart
        style_native(chart, c, self.doc.style.theme, self.latin, self.cs)
        chart.has_title = False  # the slide's title says it
        if self.doc.style.theme == "dark":
            chart.font.color.rgb = _rgb(self.t["text"])

    def list_notes_master(self) -> None:
        """python-pptx relates the notes master but doesn't list it in presentation.xml. Apple's Office importer
        (Quick Look, iPhone previews of the SMS link's download) hangs on such a deck, so list it."""
        pres = self.prs.part._element
        rid = next((r.rId for r in self.prs.part.rels.values() if r.reltype == RT.NOTES_MASTER), None)
        if rid is None or pres.find(qn("p:notesMasterIdLst")) is not None:
            return
        lst = pres.makeelement(qn("p:notesMasterIdLst"), {})
        lst.append(lst.makeelement(qn("p:notesMasterId"), {qn("r:id"): rid}))
        pres.find(qn("p:sldMasterIdLst")).addnext(lst)  # schema order: sldMasterIdLst, notesMasterIdLst, ...

    def save(self) -> bytes:
        self.list_notes_master()
        self.prs.core_properties.title = self.doc.title
        self.prs.core_properties.author = "BolDoc"
        buf = io.BytesIO()
        self.prs.save(buf)
        return buf.getvalue()


def render(doc: Document) -> bytes:
    deck = _Deck(doc)
    deck.title_slide(doc.title, subtitle(doc))
    explicit = any(isinstance(b, Slide) for b in doc.blocks)
    last_title = doc.title
    pending: List[str] = []

    def flush() -> None:
        nonlocal pending
        if pending:
            deck.bullets(last_title, pending)
        pending = []

    for b in doc.blocks:
        if isinstance(b, Slide):
            flush()
            last_title = b.title
            if b.layout == "title":
                deck.title_slide(b.title, " · ".join(b.bullets[:2]))
            elif b.layout == "section":
                deck.section(b.title, b.notes)
            elif b.layout == "two_column":
                deck.two_column(b.title, b.bullets, b.right, b.notes)
            elif b.layout == "quote":
                deck.quote(b.title, b.quote or " ".join(b.bullets[:1]) or b.title, b.notes)
            elif b.layout == "closing":
                deck.closing(b.title, b.bullets, b.notes)
            else:
                deck.bullets(b.title, b.bullets, b.notes)
        elif isinstance(b, Heading):
            flush()
            last_title = b.text
            if explicit:
                deck.section(b.text)
        elif isinstance(b, Table):
            flush()
            deck.table(b, last_title)
        elif isinstance(b, Chart):
            flush()
            deck.chart(b, last_title)
        elif isinstance(b, Paragraph) and b.text.strip():
            pending.append(b.text.strip())
        elif isinstance(b, ListBlock):
            pending.extend(b.items)
        elif isinstance(b, KeyValue):
            pending.extend(f"{p[0]}: {p[1]}" for p in b.pairs if len(p) >= 2)
        elif isinstance(b, Signature):
            pending.append(" — ".join(x for x in (b.name, b.designation) if x))
    flush()
    return deck.save()
