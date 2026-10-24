"""Text from a file the caller added for BolDoc, read in a process of its own: python-docx, openpyxl,
python-pptx and pypdf are pure Python, and a big file would hold control's interpreter (and the live calls it serves)
for seconds. The file comes on stdin; {"text", "pages", "scanned"} goes out as JSON (or {"error"}).

    python -m apps.control.extract <kind> < file        kind: pdf | docx | xlsx | pptx | csv | txt

A PDF with almost no text on its pages is a scan ("scanned": true): materials.py sends it to Sarvam's document reader.
Tables come out a row a line, cells joined by " | ", so the writer reads a sheet's figures as they stand.
"""
from __future__ import annotations

import asyncio
import csv
import io
import json
import sys
from typing import Tuple

MAX_CHARS = 200_000    # of a file's text: a 40-page report or a long sheet
MAX_ROWS = 500         # a sheet's rows read
MAX_COLS = 40
KINDS = ("pdf", "docx", "xlsx", "pptx", "csv", "txt")


def decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16") if data[:2] in (b"\xff\xfe", b"\xfe\xff") else ("utf-8-sig",):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("cp1252", errors="replace")


def _row(cells) -> str:
    vals = ["" if v is None else str(v).strip() for v in list(cells)[:MAX_COLS]]
    while vals and not vals[-1]:
        vals.pop()
    return " | ".join(vals)


def docx_text(data: bytes) -> Tuple[str, int]:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = Document(io.BytesIO(data))
    out = []
    for el in d.element.body.iterchildren():  # paragraphs and tables in the order the reader meets them
        if el.tag == qn("w:p"):
            text = Paragraph(el, d).text.strip()
            if text:
                out.append(text)
        elif el.tag == qn("w:tbl"):
            rows = [_row(c.text for c in r.cells) for r in Table(el, d).rows]
            out.append("\n".join(r for r in rows if r.strip(" |")))
    return "\n".join(out), 0


def xlsx_text(data: bytes) -> Tuple[str, int]:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)  # the figures a formula shows, not it
    out = []
    for ws in wb.worksheets:
        rows = []
        for cells in ws.iter_rows(values_only=True, max_row=MAX_ROWS):
            line = _row(cells)
            if line.strip(" |"):
                rows.append(line)
        if rows:
            out.append(f"Sheet: {ws.title}\n" + "\n".join(rows))
    return "\n\n".join(out), len(wb.worksheets)


def pptx_text(data: bytes) -> Tuple[str, int]:
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    out = []
    for i, slide in enumerate(prs.slides, start=1):
        parts = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip())
            elif getattr(shape, "has_table", False):
                parts += [_row(c.text for c in r.cells) for r in shape.table.rows]
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            parts.append("Notes: " + slide.notes_slide.notes_text_frame.text.strip())
        if parts:
            out.append(f"Slide {i}:\n" + "\n".join(parts))
    return "\n\n".join(out), len(prs.slides)


def csv_text(data: bytes) -> Tuple[str, int]:
    text = decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = [_row(r) for i, r in zip(range(MAX_ROWS), csv.reader(io.StringIO(text), dialect))]
    return "\n".join(r for r in rows if r.strip(" |")), 0


def extract(kind: str, data: bytes) -> Tuple[str, int, bool]:
    """(text, pages or sheets or slides, scanned)."""
    if kind == "pdf":
        from .pdftext import extract as pdf_extract
        _title, text, pages = pdf_extract(data, 200)
        # a scan: under ~25 characters of text a page (a page number, a stamp) is no text at all
        return text, pages, len(text.strip()) < 25 * max(1, pages)
    readers = {"docx": docx_text, "xlsx": xlsx_text, "pptx": pptx_text, "csv": csv_text}
    if kind in readers:
        text, pages = readers[kind](data)
        return text, pages, False
    if kind == "txt":
        return decode(data), 0, False
    raise ValueError(f"can't read a {kind} file")


async def extract_async(kind: str, data: bytes, timeout_s: float = 60.0) -> Tuple[str, int, bool]:
    """extract in a process of its own. Raises ValueError when the file can't be read in time."""
    proc = await asyncio.create_subprocess_exec(sys.executable, "-m", "apps.control.extract", kind,
                                                stdin=asyncio.subprocess.PIPE, stdout=asyncio.subprocess.PIPE,
                                                stderr=asyncio.subprocess.DEVNULL)
    try:
        out, _ = await asyncio.wait_for(proc.communicate(data), timeout_s)
    except asyncio.TimeoutError:
        raise ValueError("the file took too long to read") from None
    finally:
        if proc.returncode is None:
            proc.kill()
    if proc.returncode != 0:
        raise ValueError("the file couldn't be read")
    res = json.loads(out.decode("utf-8"))
    if res.get("error"):
        raise ValueError(res["error"])
    return res["text"], res["pages"], res["scanned"]


def main() -> None:
    kind = sys.argv[1] if len(sys.argv) > 1 else ""
    try:
        text, pages, scanned = extract(kind, sys.stdin.buffer.read())
        out = {"text": text[:MAX_CHARS], "pages": pages, "scanned": scanned}
    except Exception as e:  # a broken, locked or mislabelled file: said plainly, the process exits cleanly
        out = {"error": str(e)[:200] if isinstance(e, ValueError) else "the file couldn't be read (is it damaged?)"}
    sys.stdout.write(json.dumps(out, ensure_ascii=False))


if __name__ == "__main__":
    main()
