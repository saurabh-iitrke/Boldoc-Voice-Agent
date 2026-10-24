"""Formats: decks with layouts, tables and charts; workbooks with chart and summary sheets; long Word and PDF
with a cover, contents, page numbers and numbered sections; charts whose labels stay shaped text (Devanagari)."""
import io
import re
import zipfile

from docengine import render
from docengine.model import Column, Document, Table


def test_deck_layouts_tables_and_charts(pitch_deck):
    from pptx import Presentation

    prs = Presentation(io.BytesIO(render(pitch_deck, "pptx").data))
    titles = [s.shapes.title.text for s in prs.slides]
    assert len(titles) == 11
    assert titles[5:7] == ["Project cost", "Project cost (contd.)"]  # 13 rows: 10, then 3 and the total
    assert titles[8:10] == ["Long list", "Long list (contd.)"]     # 9 bullets: 6, then 3
    assert any(sh.has_chart for sh in prs.slides[7].shapes)          # a native chart, editable in PowerPoint
    assert prs.slides[1].notes_slide.notes_text_frame.text == "Open with the story"
    pres = zipfile.ZipFile(io.BytesIO(render(pitch_deck, "pptx").data)).read("ppt/presentation.xml").decode()
    assert "notesMasterIdLst" in pres                                # Apple previews hang without it
    last = prs.slides[6].shapes[-1].table                            # the total row on the last part
    assert last.cell(len(last.rows) - 1, 0).text == "Total" and last.cell(len(last.rows) - 1, 1).text == "₹9,10,000"


def test_workbook_has_chart_and_summary_sheets(project_report):
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(render(project_report, "xlsx").data))
    assert wb.sheetnames[:2] == ["Capital cost", "Profit and loss"]
    assert wb["Sales and profit"]._charts and wb["Cost split"]._charts
    summary = [[c.value for c in row] for row in wb["Summary"].iter_rows()]
    assert any(isinstance(v, str) and v.startswith("='Capital cost'!") for row in summary for v in row)


def test_long_word_document_has_cover_contents_page_numbers_and_numbered_sections(project_report):
    z = zipfile.ZipFile(io.BytesIO(render(project_report, "docx").data))
    body = z.read("word/document.xml").decode()
    assert 'TOC \\o "1-2"' in body and "1.1" not in body.split("TOC")[0]
    assert re.search(r">2\s+Project cost<", body)  # the contents' cached text
    assert ">2.1  <" in body and ">Capital items<" in body  # a heading: its number (in the accent), then its text
    assert "word/charts/chart1.xml" in z.namelist() and "word/charts/chart2.xml" in z.namelist()
    # each chart with its workbook: Word's "Edit data", and Apple's previewer hung on charts without one
    assert "word/embeddings/Microsoft_Excel_Worksheet2.xlsx" in z.namelist()
    assert b'<c:externalData r:id="rId1">' in z.read("word/charts/chart1.xml")
    footers = "".join(z.read(n).decode() for n in z.namelist() if n.startswith("word/footer"))
    assert 'w:instr="PAGE"' in footers and 'w:instr="NUMPAGES"' in footers
    assert "w:updateFields" in z.read("word/settings.xml").decode()  # Word fills the contents' page numbers


def test_long_pdf_source_has_cover_contents_and_page_counters(project_report):
    html = render(project_report, "html").data.decode()
    assert '<section class="cover">' in html and '<nav class="toc">' in html
    ids = re.findall(r'<h[23] id="([^"]+)"', html)
    assert [f"#{i}" for i in ids] == re.findall(r'<a href="(#[^"]+)"', html)  # every section in the contents
    assert "target-counter(attr(href), page)" in html and "counter(pages)" in html
    assert "@page :first" in html and "string(doctitle)" in html


def test_pdf_charts_keep_their_labels_as_text_not_svg(deck_hi):
    """WeasyPrint draws SVG text letter by letter: Devanagari inside the SVG lost its matras."""
    html = render(deck_hi, "html").data.decode()
    svgs = re.findall(r"<svg.*?</svg>", html, re.S)
    assert svgs and not any("<text" in s for s in svgs)
    assert "साल 1" in html and "₹40,00,000" in html  # categories and paise-free axis labels, as page text


def test_pdf_of_a_deck_is_one_page_per_slide(pitch_deck):
    html = render(pitch_deck, "html").data.decode()
    assert html.count('<section class="slide') == 11 and "size: 297mm 167.06mm" in html


def test_totals_are_labelled_in_the_documents_language():
    doc = Document(id="t", title="खर्च", lang="hi-IN", script="Deva", blocks=[Table(
        total_row=True, columns=[Column(name="चीज़"), Column(name="रकम", type="currency")], rows=[["होटल", 1200]])])
    assert "कुल" in render(doc, "md").data.decode() and "कुल" in render(doc, "html").data.decode()


def designed_report(doc_type="report"):
    from docengine.model import Heading, KeyValue, ListBlock, Style
    return Document(
        id="d-look", title="Project report", lang="en-IN", doc_type=doc_type,
        meta={"subtitle": "For a ₹4,00,000 MUDRA Kishor loan", "author": "Saurabh Kumar, Kumar Tea & Snacks, Jaipur",
              "date": "23 October 2026"},
        style=Style(cover=True, toc=True, page_numbers=True),
        blocks=[Heading(text="Costs"),
                KeyValue(pairs=[["Loan sought", "₹4,00,000"], ["Own contribution", "₹1,00,000"],
                                ["Shop rent", "[Your shop rent per month]"]]),
                Table(caption="Table 1.1: Detailed project cost", columns=[Column(name="Item"),
                      Column(name="Amount", type="currency")], rows=[["Machine", 120000], ["Fridge", 45000]],
                      total_row=True),
                Heading(text="Sources"),
                ListBlock(ordered=True, items=["PMMY. https://www.mudra.org.in/offerings (read 23 October 2026)"])],
        assumptions=["Fill in [Your shop rent per month]"])


def test_a_long_documents_cover_tables_and_sources_read_as_designed():
    """The cover says what the document is for, who presents it and when; amounts are whole rupees where every one is
    whole (an invoice keeps its paise); the Sources list is set as references. A fact the caller left out is a blank
    to fill in, and the writer's notes are never printed."""
    from docengine.renderers import BLANK
    html = render(designed_report(), "html").data.decode()
    assert '<p class="subtitle">For a ₹4,00,000 MUDRA Kishor loan</p>' in html
    assert '<p class="author">Saurabh Kumar, Kumar Tea &amp; Snacks, Jaipur</p>' in html
    assert '<p class="date">23 October 2026</p>' in html and "Date:" not in html  # the date sits in the byline
    assert "<td class=num>₹1,65,000</td>" in html and "₹1,20,000.00" not in html
    assert 'class="sources"' in html and 'table class="kv"' in html
    assert "Assumptions" not in html and "[Your shop rent per month]" not in html and BLANK in html
    assert "<td class=num>₹1,65,000.00</td>" in render(designed_report("invoice"), "html").data.decode()
    doc = designed_report()
    assert doc.assumptions and "[Your shop rent per month]" in doc.all_text()  # the stored document keeps both


def test_a_long_word_document_is_a4_with_the_same_cover_and_clean_tables():
    from docx import Document as Docx
    data = render(designed_report(), "docx").data
    d = Docx(io.BytesIO(data))
    sec = d.sections[0]
    assert (round(sec.page_width.mm), round(sec.page_height.mm)) == (210, 297)  # A4, not python-docx's US Letter
    text = [p.text for p in d.paragraphs]
    assert text[:2] == ["Project report", "For a ₹4,00,000 MUDRA Kishor loan"]
    assert "Saurabh Kumar, Kumar Tea & Snacks, Jaipur" in text and "23 October 2026" in text
    assert not any("Assumptions" in x or "[Your shop rent" in x for x in text)
    assert any("______________" in c.text for tbl in d.tables for row in tbl.rows for c in row.cells)
    body = zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml").decode()
    assert "Table Grid" not in body and '<w:insideV w:val="nil"/>' in body  # ruled rows, no grid
    assert [c.text for c in d.tables[1].rows[-1].cells] == ["Total", "₹1,65,000"]
    assert [c.text for c in d.tables[0].rows[0].cells] == ["Loan sought", "₹4,00,000"]  # key figures as a table


def test_workbook_tabs_are_named_by_what_they_hold():
    from openpyxl import load_workbook

    from docengine.model import Style
    doc = designed_report().model_copy(update={"style": Style()})
    long = Table(caption="Table 2.1: Three-year financial projections for the expanded cafe",
                 columns=[Column(name="Year"), Column(name="Sales", type="currency")], rows=[["1", 300000]])
    wb = load_workbook(io.BytesIO(render(doc.model_copy(update={"blocks": [*doc.blocks, long]}), "xlsx").data))
    assert wb.sheetnames[:2] == ["Detailed project cost", "Three-year financial"]
    assert wb["Detailed project cost"]["B2"].number_format.endswith('"₹"#,##0')  # whole rupees
