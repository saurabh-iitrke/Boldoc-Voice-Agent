import io

import pytest

from docengine import render, validate
from docengine.renderers import FORMATS, RendererUnavailable, ascii_slug, content_disposition
from docengine.validators import find_font_file, missing_glyphs

LOCAL_FORMATS = [f for f in FORMATS if f != "pdf"]


@pytest.mark.parametrize("fmt", LOCAL_FORMATS)
def test_every_format_renders_and_validates(sample_doc, fmt):
    out = render(sample_doc, fmt)
    assert out.data and out.filename.endswith(f"_v{sample_doc.version}.{fmt}")
    assert validate(sample_doc, out) == []


def test_pdf_renders_when_weasyprint_available(sample_doc):
    try:
        out = render(sample_doc, "pdf")
    except RendererUnavailable:
        pytest.skip("WeasyPrint/Pango not installed (runs on the Linux server)")
    assert validate(sample_doc, out) == []


def test_docx_complex_script_properties(letter_hi):
    from docx import Document as Docx
    from docx.oxml.ns import qn

    d = Docx(io.BytesIO(render(letter_hi, "docx").data))
    deva_runs = [r for p in d.paragraphs for r in p.runs if any(ord(c) >= 0x0900 for c in r.text)]
    assert deva_runs
    for r in deva_runs:
        assert r._element.rPr.find(qn("w:rFonts")).get(qn("w:cs")) == "Nirmala UI"
    bold = [r for r in deva_runs if r.bold]
    assert bold and all(r._element.rPr.find(qn("w:bCs")) is not None for r in bold)


def test_xlsx_formulas_and_formats(expenses):
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(render(expenses, "xlsx").data))
    ws = wb["Expenses"]
    assert ws["B5"].value == "=SUM(B2:B4)"
    assert "₹" in ws["B2"].number_format
    assert ws["B4"].value == 150000  # "1.5 lakh" parsed to a real number
    assert ws.freeze_panes == "A2"
    assert "Notes" in wb.sheetnames  # heading + key-value prose kept, not dropped


def test_csv_has_bom_and_raw_values(expenses):
    data = render(expenses, "csv").data
    assert data.startswith("﻿".encode())
    text = data.decode("utf-8-sig").splitlines()
    assert text[0] == "Item,Amount,Date" and text[1].startswith("Hotel,12500")


def test_markdown_table_alignment(expenses):
    md = render(expenses, "md").data.decode()
    assert "|---|---:|---|" in md and "| ₹1,64,840 |" in md  # total row: Indian grouping, whole rupees


def test_rtl_document(urdu):
    html = render(urdu, "html").data.decode()
    assert 'dir="rtl"' in html and 'lang="ur-IN"' in html


def test_pptx_slides(letter_en, sample_doc):
    from pptx import Presentation

    prs = Presentation(io.BytesIO(render(letter_en, "pptx").data))
    assert len(prs.slides) >= 2  # title slide + derived content slide(s)


def test_filenames_are_ascii_safe():
    assert ascii_slug("छुट्टी के लिए आवेदन") == "document"
    assert ascii_slug("Leave Application / छुट्टी") == "Leave_Application"
    cd = content_disposition("छुट्टी आवेदन", "pdf", 3)
    assert 'filename="document_v3.pdf"' in cd and "filename*=UTF-8''" in cd


def test_unsupported_format(letter_en):
    with pytest.raises(ValueError):
        render(letter_en, "pages")


def test_devanagari_glyph_coverage(letter_hi):
    font = find_font_file("Noto Sans Devanagari")
    if font is None:
        pytest.skip("Noto Sans Devanagari not installed (installed on the server via fonts-noto-core)")
    assert missing_glyphs(letter_hi.all_text(), font) == []


def test_order_validator_catches_out_of_order_rpr():
    from docx import Document as Docx
    from docx.oxml import OxmlElement
    from docx.shared import Pt

    from docengine.validators import _ooxml_order_problems

    d = Docx()
    r = d.add_paragraph().add_run("x")
    r.font.size = Pt(11)
    r._element.get_or_add_rPr().append(OxmlElement("w:bCs"))  # after w:sz: invalid for Word
    buf = io.BytesIO()
    d.save(buf)
    assert _ooxml_order_problems(buf.getvalue())


def test_content_disposition_escapes_slash():
    cd = content_disposition("Q3 / Q4 plan", "docx", 1)
    assert "%2F" in cd and "/ Q4" not in cd.split("filename*=")[1]


def test_title_slide_without_subtitle_has_no_empty_placeholder(sample_doc):
    from pptx import Presentation

    sample_doc.meta = {}
    prs = Presentation(io.BytesIO(render(sample_doc, "pptx").data))
    # an empty placeholder shows "Click to add subtitle" while editing; the accent band is a shape, not a placeholder
    assert all(ph.has_text_frame and ph.text_frame.text for ph in prs.slides[0].placeholders)


def test_pdf_validator_flags_non_embedded_font(letter_en):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject

    w = PdfWriter()
    page = w.add_blank_page(200, 200)
    helvetica = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                  NameObject("/Subtype"): NameObject("/Type1"),
                                  NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject(
        {NameObject("/F1"): w._add_object(helvetica)})})
    buf = io.BytesIO()
    w.write(buf)
    from docengine.renderers import RenderedFile

    problems = validate(letter_en, RenderedFile(buf.getvalue(), "pdf", "application/pdf", "x.pdf"))
    assert problems == ["pdf: font /Helvetica not embedded"]
