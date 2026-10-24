"""Reading the files a caller adds: each kind's text, tables a row a line, a scan told apart, a
mislabelled file turned away, and the reading done in a process of its own."""
import asyncio
import io

import pytest

from apps.control import extract
from apps.control.materials import MaterialError, kind_of
from tests.control.test_research import tiny_pdf

pytest.importorskip("pypdf")


def test_word_excel_powerpoint_csv_and_text_come_out_as_text_with_tables_row_by_row():
    from docx import Document
    d = Document()
    d.add_paragraph("Quotation for Kumar Tea & Snacks")
    t = d.add_table(rows=2, cols=2)
    for r, (a, b) in enumerate([("Item", "Amount"), ("Coffee machine", "1,20,000")]):
        t.cell(r, 0).text, t.cell(r, 1).text = a, b
    d.add_paragraph("Valid till 30 November 2026.")
    buf = io.BytesIO()
    d.save(buf)
    text, _, scanned = extract.extract("docx", buf.getvalue())
    assert text.splitlines() == ["Quotation for Kumar Tea & Snacks", "Item | Amount", "Coffee machine | 1,20,000",
                                 "Valid till 30 November 2026."] and not scanned

    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Costs"
    for row in (["Item", "Amount"], ["Coffee machine", 120000], [None, None], ["Fridge", 45000]):
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    text, sheets, _ = extract.extract("xlsx", buf.getvalue())
    assert text == "Sheet: Costs\nItem | Amount\nCoffee machine | 120000\nFridge | 45000" and sheets == 1

    from pptx import Presentation
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Traction"
    s.placeholders[1].text = "1,200 cups a day"
    s.notes_slide.notes_text_frame.text = "Mention the festival season"
    buf = io.BytesIO()
    prs.save(buf)
    text, slides, _ = extract.extract("pptx", buf.getvalue())
    assert text == "Slide 1:\nTraction\n1,200 cups a day\nNotes: Mention the festival season" and slides == 1

    assert extract.extract("csv", b"Month;Sales\nApril;2,40,000\nMay;2,55,000\n")[0] == \
        "Month | Sales\nApril | 2,40,000\nMay | 2,55,000"
    assert extract.extract("txt", "मासिक बिक्री ₹2,40,000".encode("utf-16"))[0] == "मासिक बिक्री ₹2,40,000"


def test_a_pdf_with_text_is_read_and_one_without_is_a_scan():
    text, pages, scanned = extract.extract("pdf", tiny_pdf("Project cost Rs 5,00,000 for the cafe upgrade in Jaipur."))
    assert "Project cost Rs 5,00,000" in text and pages == 1 and not scanned
    assert extract.extract("pdf", tiny_pdf(""))[2] is True  # a page with no text: a scan, for the reader


def test_files_are_read_in_a_process_of_their_own_and_a_broken_one_says_so():
    text, _, _ = asyncio.run(extract.extract_async("csv", b"a,b\n1,2\n"))
    assert text == "a | b\n1 | 2"
    with pytest.raises(ValueError, match="couldn't be read"):
        asyncio.run(extract.extract_async("docx", b"PK this is not a Word file"))


def test_a_file_is_taken_by_its_name_and_checked_by_its_first_bytes():
    assert kind_of("Quote.PDF", b"%PDF-1.4 ...") == ("pdf", "application/pdf")
    assert kind_of("photo.jpeg", b"\xff\xd8\xff\xe0...")[0] == "image"
    for name, data in (("report.pdf", b"PK\x03\x04"), ("notes.exe", b"MZ"), ("scan.png", b"%PDF-")):
        with pytest.raises(MaterialError) as e:
            kind_of(name, data)
        assert e.value.status == 415
